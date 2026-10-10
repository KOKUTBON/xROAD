"""Reporting and executive-dashboard product service."""

from __future__ import annotations

import csv
from datetime import UTC, datetime
import io
import json
from pathlib import Path
import textwrap
from typing import Any
from uuid import uuid4

from apps.api.db import repositories
from apps.api.services import artifact_storage, intervention_portfolio
from apps.api.settings import ApiSettings
from apps.api.storage import resolve_artifact_download
from road_risk_modeling.data_governance import publishable_fields_for_role
from sqlalchemy.exc import SQLAlchemyError


# Default roles used by the matrix-projected report serializers below.
# The public listing surface uses `"public"`; the authenticated org
# history surface uses `"analyst"`. Routes can pass other roles
# (`data_operator`, `admin`) to broaden visibility (e.g. to expose
# `file_storage_uri` to ops tooling).
REPORT_PUBLIC_ROLE = "public"
REPORT_DEFAULT_AUTHENTICATED_ROLE = "analyst"

# Strict precedence (most-privileged first) used by
# `resolve_report_role_for_auth` to pick the matrix role for a given
# caller. The first role on this list that the caller holds wins; if
# they hold none, the surface falls back to `"public"`.
_REPORT_ROLE_PRECEDENCE: tuple[str, ...] = (
    "admin",
    "data_operator",
    "analyst",
    "municipality_viewer",
)


def resolve_report_role_for_auth(role_names: set[str] | frozenset[str]) -> str:
    """Pick the matrix role to use for an authenticated report-reading caller.

    Walks `_REPORT_ROLE_PRECEDENCE` and returns the first role the
    caller holds; if they hold none of those four, returns `"public"`
    (the matrix-visible fields are then the union of `public`-visible
    entries). Centralized here so routes don't duplicate the mapping.
    """
    held = set(role_names)
    for candidate in _REPORT_ROLE_PRECEDENCE:
        if candidate in held:
            return candidate
    return REPORT_PUBLIC_ROLE


REPORTING_OVERVIEW_CONTRACT_VERSION = "reporting_overview_v1"
REPORT_HISTORY_CONTRACT_VERSION = "report_history_v1"
REPORT_DETAIL_CONTRACT_VERSION = "report_detail_v1"
REPORT_GENERATION_CONTRACT_VERSION = "report_generation_v1"
REPORT_LIFECYCLE_CONTRACT_VERSION = "report_lifecycle_v1"
PUBLIC_REPORTS_CONTRACT_VERSION = "public_reports_v1"
EXECUTIVE_OVERVIEW_CONTRACT_VERSION = "executive_overview_v1"
REPORT_GENERATION_TASK_NAME = "services.workers.tasks.reports.run_report"

DEFAULT_HISTORY_LIMIT = 50
DEFAULT_SEGMENT_EXPORT_LIMIT = 25
DEFAULT_TILE_EXPORT_LIMIT = 12
REPORT_LIFECYCLE_STATUSES = (
    "draft",
    "queued",
    "running",
    "generated",
    "reviewed",
    "published",
    "archived",
    "failed",
)
REPORT_DOWNLOAD_READY_STATUSES = frozenset({"generated", "reviewed", "published"})
PUBLIC_REPORT_TYPES = frozenset({"public_transparency"})
PUBLICATION_TIERS = frozenset({"restricted", "internal", "partner_restricted", "analyst_only", "public"})


class ReportingHubError(RuntimeError):
    """Raised when the reporting service cannot build a product response."""


class ReportingHubAuthorizationError(ReportingHubError):
    """Raised when a report action is blocked by publication or role policy."""


REPORT_TEMPLATES: dict[str, dict[str, Any]] = {
    "executive_overview": {
        "label": "Executive Safety Overview",
        "audience": "executive",
        "description": (
            "High-level safety KPIs, priority areas, and expected portfolio impact for "
            "directors, mayors, and mobility executives."
        ),
        "required_role": "municipality_viewer",
        "supported_formats": ["pdf", "md", "xlsx"],
        "default_format": "pdf",
    },
    "technical_hotspot": {
        "label": "Technical Hotspot Pack",
        "audience": "analyst",
        "description": (
            "Analyst-oriented corridor, segment, and hotspot export for technical review, "
            "handoff, and follow-on GIS work."
        ),
        "required_role": "analyst",
        "supported_formats": ["pdf", "md", "csv", "geojson", "xlsx"],
        "default_format": "csv",
    },
    "intervention_portfolio": {
        "label": "Intervention Portfolio Export",
        "audience": "analyst",
        "description": (
            "Managed municipal intervention portfolio with status, owner, cost class, "
            "expected benefit, priority score, and monitoring placeholders."
        ),
        "required_role": "analyst",
        "supported_formats": ["pdf", "md", "csv", "xlsx"],
        "default_format": "csv",
    },
    "data_quality_governance": {
        "label": "Data Quality & Governance Report",
        "audience": "data_governance",
        "description": (
            "Data coverage, source freshness, governance posture, publication tier, "
            "and export-readiness evidence for operators, admins, and legal review."
        ),
        "required_role": "data_operator",
        "supported_formats": ["pdf", "md", "csv", "xlsx"],
        "default_format": "xlsx",
    },
    "public_transparency": {
        "label": "Public Transparency Report",
        "audience": "public",
        "description": (
            "Citizen-safe city summary with aggregated KPIs, public corridors, "
            "public score tiles, methodology context, and no restricted crash records."
        ),
        "required_role": "municipality_viewer",
        "supported_formats": ["pdf", "md", "csv", "geojson", "xlsx"],
        "default_format": "pdf",
    },
}


def _resolve_path(path_value: Path | str, *, root_dir: Path) -> Path:
    candidate = Path(path_value)
    if candidate.is_absolute():
        return candidate
    return root_dir / candidate


def _load_json(path: Path) -> dict[str, Any]:
    if not path.exists():
        raise ReportingHubError(f"Missing reporting artifact: {path}")
    try:
        return json.loads(path.read_text(encoding="utf-8"))
    except json.JSONDecodeError as exc:
        raise ReportingHubError(f"Invalid JSON artifact: {path}") from exc


def _load_csv_rows(path: Path, *, limit: int | None = None) -> list[dict[str, str]]:
    if not path.exists():
        return []
    with path.open("r", encoding="utf-8", newline="") as handle:
        reader = csv.DictReader(handle)
        rows = list(reader)
    return rows[:limit] if limit is not None else rows


def _load_geojson(path: Path) -> dict[str, Any]:
    payload = _load_json(path)
    if payload.get("type") != "FeatureCollection":
        raise ReportingHubError(f"Expected GeoJSON FeatureCollection at: {path}")
    return payload


def _published_runs(summary: dict[str, Any]) -> list[dict[str, Any]]:
    runs = summary.get("published_runs")
    return runs if isinstance(runs, list) else []


def _find_run(summary: dict[str, Any], city_key: str) -> dict[str, Any]:
    for run in _published_runs(summary):
        if str(run.get("city_key")) == city_key:
            return run
    raise ReportingHubError(f"Unknown city_key: {city_key}")


def _city_catalog(summary: dict[str, Any]) -> list[dict[str, Any]]:
    return sorted(
        [
            {
                "city_key": str(run.get("city_key") or ""),
                "display_name": str(run.get("display_name") or run.get("city_key") or ""),
                "prediction_run_label": str(run.get("prediction_run_label") or ""),
                "occurred_years": list(run.get("occurred_years") or []),
            }
            for run in _published_runs(summary)
            if run.get("city_key")
        ],
        key=lambda row: row["display_name"],
    )


def _first_city_key(summary: dict[str, Any]) -> str:
    catalog = _city_catalog(summary)
    if not catalog:
        raise ReportingHubError("No published cities are available for reporting.")
    return str(catalog[0]["city_key"])


def _load_dashboard_summary(settings: ApiSettings) -> dict[str, Any]:
    return _load_json(_resolve_path(settings.dashboard_summary_path, root_dir=settings.root_dir))


def _find_output_path(entries: list[dict[str, Any]], suffix: str) -> str | None:
    for entry in entries:
        path = entry.get("path")
        if isinstance(path, str) and path.endswith(suffix):
            return path
    return None


def _intervention_brief_path(settings: ApiSettings, run: dict[str, Any]) -> Path:
    prediction_run_label = str(run.get("prediction_run_label") or "").strip()
    if not prediction_run_label:
        raise ReportingHubError("The selected city does not expose a prediction_run_label.")
    return _resolve_path(
        Path("data") / "intervention_briefs" / prediction_run_label / "intervention_brief.json",
        root_dir=settings.root_dir,
    )


def _load_intervention_brief(settings: ApiSettings, run: dict[str, Any]) -> dict[str, Any]:
    return _load_json(_intervention_brief_path(settings, run))


def _confidence_label(transferability: dict[str, Any]) -> str:
    gap = max(
        float(transferability.get("cold_start_gap_vs_temporal") or 0.0),
        float(transferability.get("history_free_gap_vs_temporal") or 0.0),
    )
    if gap <= 0.05:
        return "high"
    if gap <= 0.10:
        return "medium"
    return "watch"


def _full_report_record(report: Any) -> dict[str, Any]:
    """Build the un-filtered, fully-populated report record dict.

    This is the canonical shape covering every matrix-classified report
    field. Callers that need the unfiltered record (e.g. report
    generation / lifecycle internals, or admin tooling) can use this
    directly; the public-facing surfaces should always go through
    `_serialize_report_for_role` instead so they inherit the matrix gate.
    """
    parameters = dict(getattr(report, "parameters_json", {}) or {})
    summary = dict(getattr(report, "summary_json", {}) or {})
    status = str(getattr(report, "status", "draft") or "draft")
    file_storage_uri = getattr(report, "file_storage_uri", None)
    visibility_tier = str(getattr(report, "visibility_tier", "restricted") or "restricted")
    return {
        "report_id": str(getattr(report, "id", "")),
        "organization_id": str(getattr(report, "organization_id", "")),
        "title": str(getattr(report, "title", "")),
        "report_type": str(getattr(report, "report_type", "")),
        "status": status,
        "visibility_tier": visibility_tier,
        "output_format": str(getattr(report, "output_format", "pdf")),
        "city_key": parameters.get("city_key"),
        "audience": parameters.get("audience"),
        "summary": summary,
        "review_state": summary.get("review_state") or _review_state_for_status(status),
        "publication_state": _publication_state(status, visibility_tier),
        "download_state": "ready" if file_storage_uri and status in REPORT_DOWNLOAD_READY_STATUSES else "pending",
        "lifecycle_actions": _available_lifecycle_actions(report),
        "file_storage_uri": file_storage_uri,
        "created_at": getattr(getattr(report, "created_at", None), "isoformat", lambda: None)(),
        "generated_at": getattr(getattr(report, "generated_at", None), "isoformat", lambda: None)(),
        "published_at": getattr(getattr(report, "published_at", None), "isoformat", lambda: None)(),
    }


def _serialize_report_for_role(report: Any, role: str) -> dict[str, Any]:
    """Project a report record through the matrix for `role`.

    Every key on the full record is checked against
    `publishable_fields_for_role("report", role)`. Unregistered keys
    are default-denied; classified keys whose `role_visibility` does
    not include `role` are also default-denied.

    The result is what any role-gated caller should hand back to the
    client; pair it with `_report_redaction_policy_meta(role)` for the
    response audit block.
    """
    publishable = set(publishable_fields_for_role("report", role))
    full = _full_report_record(report)
    return {key: value for key, value in full.items() if key in publishable}


def _serialize_report(report: Any) -> dict[str, Any]:
    """Back-compat wrapper used by the authenticated read paths.

    Defaults to the analyst projection. Pre-matrix callers that needed
    the full unfiltered record now get the analyst subset; the small
    behavioral tightening this introduces (file_storage_uri is no
    longer surfaced to analyst callers via `_serialize_report`) matches
    the matrix's existing `restricted` classification for
    `file_storage_uri`. The download endpoint computes its own access
    block from the raw URI before it ever calls back through the
    matrix filter.
    """
    return _serialize_report_for_role(report, REPORT_DEFAULT_AUTHENTICATED_ROLE)


def _public_report_row(report: Any) -> dict[str, Any]:
    """Public-listing projection of a report record.

    The matrix-driven equivalent of the previous hand-rolled
    `_serialize_report` -> pop(file_storage_uri) -> recompute(download_state)
    pattern. Crucially, fields like `organization_id`, `audience`,
    `summary`, `review_state`, `lifecycle_actions`, and `created_at`
    are now default-denied for the public surface - they have no
    `public` role-visibility in the matrix.

    The repository layer pre-filters public reports to
    `status="published"` + `visibility_tier="public"`, so the
    `download_state` derived inside `_full_report_record` is already
    "ready" for any report reaching this projection (no recompute
    needed).
    """
    return _serialize_report_for_role(report, REPORT_PUBLIC_ROLE)


def _report_redaction_policy_meta(role: str) -> dict[str, Any]:
    """Return the redaction-policy descriptor attached to report responses.

    Mirrors the `redaction_policy` block emitted by
    `published_queries.query_published_resource` and
    `scenario_planner.build_*`. Lets the client / audit trail confirm
    which matrix entry governed the response.
    """
    return {
        "record_kind": "report",
        "role": role,
        "publishable_fields": sorted(publishable_fields_for_role("report", role)),
    }


def _review_state_for_status(status: str) -> str:
    if status in {"generated", "queued", "running"}:
        return "pending_review"
    if status == "reviewed":
        return "reviewed"
    if status == "published":
        return "published"
    if status == "archived":
        return "archived"
    if status == "failed":
        return "failed"
    return "draft"


def _publication_state(status: str, visibility_tier: str) -> str:
    if status == "published" and visibility_tier == "public":
        return "public"
    if status == "published":
        return "published_internal"
    if status == "archived":
        return "archived"
    return "not_published"


def _available_lifecycle_actions(report: Any) -> list[str]:
    status = str(getattr(report, "status", "draft") or "draft")
    report_type = str(getattr(report, "report_type", "") or "")
    file_storage_uri = getattr(report, "file_storage_uri", None)
    if status == "generated":
        return ["review", "archive"] if file_storage_uri else ["archive"]
    if status == "reviewed":
        actions = ["archive"]
        if report_type in PUBLIC_REPORT_TYPES and file_storage_uri:
            actions.insert(0, "publish")
        return actions
    if status == "published":
        return ["unpublish", "archive"]
    if status in {"failed", "archived"}:
        return ["reopen"]
    return []


def _report_templates() -> list[dict[str, Any]]:
    return [
        {"report_type": report_type, **template}
        for report_type, template in REPORT_TEMPLATES.items()
    ]


def _segment_geojson_path(settings: ApiSettings, run: dict[str, Any]) -> Path:
    segment_scores = run.get("segment_scores") or {}
    public_segments_path = segment_scores.get("public_segments_path")
    if isinstance(public_segments_path, str):
        return _resolve_path(public_segments_path, root_dir=settings.root_dir)
    trace = segment_scores.get("lineage_trace") or {}
    entries = trace.get("direct_publishable_outputs") or []
    path = _find_output_path(entries, ".geojson")
    if path:
        return _resolve_path(path, root_dir=settings.root_dir)
    raise ReportingHubError("No segment GeoJSON export is available for the selected city.")


def _segment_csv_path(settings: ApiSettings, run: dict[str, Any]) -> Path:
    segment_scores = run.get("segment_scores") or {}
    trace = segment_scores.get("lineage_trace") or {}
    entries = trace.get("direct_publishable_outputs") or []
    path = _find_output_path(entries, ".csv")
    if path:
        return _resolve_path(path, root_dir=settings.root_dir)
    raise ReportingHubError("No analyst CSV export is available for the selected city.")


def _tile_geojson_path(settings: ApiSettings, run: dict[str, Any]) -> Path:
    score_tiles = run.get("score_tiles") or {}
    trace = score_tiles.get("lineage_trace") or {}
    entries = trace.get("direct_publishable_outputs") or []
    path = _find_output_path(entries, ".geojson")
    if path:
        return _resolve_path(path, root_dir=settings.root_dir)
    raise ReportingHubError("No score-tile GeoJSON export is available for the selected city.")


def _top_corridor_rows(brief: dict[str, Any]) -> list[dict[str, Any]]:
    shortlist = brief.get("shortlist") or {}
    corridors = shortlist.get("top_corridors") or []
    rows: list[dict[str, Any]] = []
    for row in corridors[:5]:
        top_segments = row.get("top_segments") or []
        first_segment = top_segments[0] if top_segments else {}
        rows.append(
            {
                "corridor_label": row.get("corridor_label"),
                "predicted_crash_total": row.get("predicted_crash_total"),
                "max_public_score_index": row.get("max_public_score_index"),
                "segment_count": row.get("segment_count"),
                "lead_segment": first_segment.get("segment_label") or first_segment.get("name"),
            }
        )
    return rows


def _priority_rows(brief: dict[str, Any]) -> list[dict[str, Any]]:
    actions = brief.get("recommended_actions") or []
    rows: list[dict[str, Any]] = []
    for row in actions[:5]:
        rows.append(
            {
                "priority": row.get("priority"),
                "title": row.get("title"),
                "rationale": row.get("rationale"),
                "code": row.get("code"),
            }
        )
    return rows


def _empty_portfolio_export(organization_id: str | None, city_key: str) -> dict[str, Any]:
    return {
        "contract_version": intervention_portfolio.INTERVENTION_PORTFOLIO_EXPORT_CONTRACT_VERSION,
        "organization_id": organization_id,
        "city_key": city_key,
        "availability": "degraded",
        "unavailable_reason": "database_unavailable",
        "summary": {
            "portfolio_count": 0,
            "expected_benefit": 0.0,
            "by_status": {},
        },
        "rows": [],
    }


def _signal_summary(settings: ApiSettings, city_key: str) -> dict[str, Any]:
    try:
        from apps.api.services import vru_signals
    except Exception as exc:  # pragma: no cover - import failure is unexpected
        raise ReportingHubError("VRU/signal service is unavailable.") from exc

    try:
        overview = vru_signals.build_signal_phase_overview(settings, city_key=city_key)
        vru_summary = vru_signals.build_vru_summary(settings, city_key=city_key)
    except Exception:
        return {
            "availability": "unavailable",
            "movements_with_provider_data": None,
            "ready_queue": None,
            "vru_crossings": None,
        }
    return {
        "availability": "available",
        "movements_with_provider_data": overview.get("coverage", {}).get("summary", {}).get("movements_with_provider_data"),
        "ready_queue": overview.get("operator_board", {}).get("totals", {}).get("ready_count"),
        "vru_crossings": vru_summary.get("stats", {}).get("crossing_feature_count"),
    }


def build_executive_overview(settings: ApiSettings, *, city_key: str) -> dict[str, Any]:
    summary = _load_dashboard_summary(settings)
    run = _find_run(summary, city_key)
    brief = _load_intervention_brief(settings, run)

    overview = brief.get("overview") or {}
    transferability = brief.get("transferability") or {}
    scenario = (brief.get("scenario_countermeasures") or {}).get("recommended_summary") or {}
    recommended_portfolio = scenario.get("recommended_portfolio") or {}
    recommended_scenario = scenario.get("recommended_scenario") or {}

    return {
        "contract_version": EXECUTIVE_OVERVIEW_CONTRACT_VERSION,
        "city": {
            "city_key": city_key,
            "display_name": brief.get("display_name") or run.get("display_name") or city_key.title(),
            "prediction_run_label": run.get("prediction_run_label"),
        },
        "kpis": {
            "predicted_crash_total": overview.get("predicted_crash_total"),
            "observed_crash_total": overview.get("observed_crash_total"),
            "very_high_risk_tiles": overview.get("tiles_very_high"),
            "high_risk_corridor_count": len((brief.get("shortlist") or {}).get("top_corridors") or []),
            "recommended_action_count": len(brief.get("recommended_actions") or []),
            "expected_portfolio_reduction": recommended_portfolio.get("expected_crash_reduction"),
            "signal_coverage_movements": _signal_summary(settings, city_key).get("movements_with_provider_data"),
            "model_confidence": _confidence_label(transferability),
        },
        "budget_impact": {
            "recommended_portfolio_title": recommended_portfolio.get("title"),
            "expected_crash_reduction": recommended_portfolio.get("expected_crash_reduction"),
            "expected_relative_reduction": recommended_portfolio.get("expected_relative_reduction"),
            "recommended_scenario_title": recommended_scenario.get("title"),
            "recommended_scenario_reduction": recommended_scenario.get("expected_crash_reduction"),
        },
        "priority_areas": {
            "top_corridors": _top_corridor_rows(brief),
            "top_actions": _priority_rows(brief),
        },
        "reporting_context": {
            "transferability": transferability,
            "signal_summary": _signal_summary(settings, city_key),
        },
    }


def list_report_history(
    session: Any,
    *,
    organization_id: str,
    city_key: str | None = None,
    limit: int = DEFAULT_HISTORY_LIMIT,
    role: str = REPORT_DEFAULT_AUTHENTICATED_ROLE,
) -> dict[str, Any]:
    """List an organization's report history, matrix-projected for `role`.

    `role` selects which `report` matrix entries are publishable on the
    history rows. The default is `"analyst"`; pass `"data_operator"`
    or `"admin"` to also surface `file_storage_uri`. The response
    carries a `redaction_policy` audit block.
    """
    try:
        reports = repositories.list_reports_for_organization(session, organization_id)
        availability = "available"
        unavailable_reason = None
    except SQLAlchemyError:
        reports = []
        availability = "degraded"
        unavailable_reason = "database_unavailable"
    rows = [_serialize_report_for_role(report, role) for report in reports]
    if city_key:
        rows = [row for row in rows if row.get("city_key") == city_key]
    return {
        "contract_version": REPORT_HISTORY_CONTRACT_VERSION,
        "organization_id": organization_id,
        "filters": {"city_key": city_key, "limit": limit},
        "availability": availability,
        "unavailable_reason": unavailable_reason,
        "rows": rows[:limit],
        "redaction_policy": _report_redaction_policy_meta(role),
    }


def build_reporting_overview(
    session: Any,
    settings: ApiSettings,
    *,
    organization_id: str,
    city_key: str | None = None,
    history_limit: int = DEFAULT_HISTORY_LIMIT,
) -> dict[str, Any]:
    summary = _load_dashboard_summary(settings)
    resolved_city_key = city_key or _first_city_key(summary)
    return {
        "contract_version": REPORTING_OVERVIEW_CONTRACT_VERSION,
        "organization_id": organization_id,
        "city_catalog": _city_catalog(summary),
        "selected_city_key": resolved_city_key,
        "executive_overview": build_executive_overview(settings, city_key=resolved_city_key),
        "report_templates": _report_templates(),
        "report_history": list_report_history(
            session,
            organization_id=organization_id,
            city_key=resolved_city_key,
            limit=history_limit,
        ),
    }


def _default_title(report_type: str, display_name: str) -> str:
    if report_type == "executive_overview":
        return f"{display_name} Executive Safety Overview"
    if report_type == "technical_hotspot":
        return f"{display_name} Technical Hotspot Pack"
    if report_type == "intervention_portfolio":
        return f"{display_name} Intervention Portfolio"
    if report_type == "data_quality_governance":
        return f"{display_name} Data Quality & Governance Report"
    if report_type == "public_transparency":
        return f"{display_name} Public Transparency Report"
    return f"{display_name} Report"


def required_role_for_report(report_type: str) -> str:
    template = REPORT_TEMPLATES.get(report_type)
    if template is None:
        raise ReportingHubError(f"Unsupported report_type: {report_type}")
    return str(template["required_role"])


def _validate_report_request(report_type: str, output_format: str) -> dict[str, Any]:
    template = REPORT_TEMPLATES.get(report_type)
    if template is None:
        raise ReportingHubError(f"Unsupported report_type: {report_type}")
    if output_format not in template["supported_formats"]:
        raise ReportingHubError(
            f"Unsupported output_format '{output_format}' for report_type '{report_type}'."
        )
    return template


def _validate_visibility(report_type: str, visibility_tier: str) -> str:
    tier = (visibility_tier or "restricted").strip().lower()
    if tier not in PUBLICATION_TIERS:
        raise ReportingHubError(f"Unsupported visibility_tier: {visibility_tier}")
    if tier == "public" and report_type not in PUBLIC_REPORT_TYPES:
        raise ReportingHubError("Only public transparency reports can use public visibility.")
    return tier


def _transient_report_payload(
    *,
    organization_id: str,
    city_key: str,
    report_type: str,
    output_format: str,
    visibility_tier: str,
    title: str,
    audience: str,
    display_name: str,
) -> dict[str, Any]:
    timestamp = datetime.now(UTC).isoformat()
    return {
        "report_id": f"transient-{uuid4().hex}",
        "organization_id": organization_id,
        "title": title,
        "report_type": report_type,
        "status": "generated",
        "visibility_tier": visibility_tier,
        "output_format": output_format,
        "city_key": city_key,
        "audience": audience,
        "summary": {
            "display_name": display_name,
            "delivery_mode": "transient",
            "unavailable_reason": "database_unavailable",
        },
        "download_state": "not_persisted",
        "file_storage_uri": None,
        "created_at": timestamp,
        "generated_at": timestamp,
    }


def create_report_request(
    session: Any,
    settings: ApiSettings,
    *,
    organization_id: str,
    city_key: str,
    report_type: str,
    output_format: str,
    visibility_tier: str = "restricted",
    title: str | None = None,
) -> dict[str, Any]:
    template = _validate_report_request(report_type, output_format)
    visibility_tier = _validate_visibility(report_type, visibility_tier)
    summary = _load_dashboard_summary(settings)
    run = _find_run(summary, city_key)
    display_name = str(run.get("display_name") or city_key.title())
    report = repositories.create_report(
        session,
        organization_id=organization_id,
        report_type=report_type,
        title=title or _default_title(report_type, display_name),
        status="queued",
        visibility_tier=visibility_tier,
        output_format=output_format,
        parameters={
            "city_key": city_key,
            "prediction_run_label": run.get("prediction_run_label"),
            "audience": template["audience"],
            "supported_formats": list(template["supported_formats"]),
        },
        summary={"display_name": display_name},
    )
    repositories.record_audit_log(
        session,
        organization_id=organization_id,
        action="report.requested",
        target_type="report",
        target_id=str(report.id),
        status="queued",
        metadata={
            "report_type": report_type,
            "output_format": output_format,
            "city_key": city_key,
        },
    )
    return {
        "contract_version": REPORT_GENERATION_CONTRACT_VERSION,
        "report": _serialize_report(report),
    }


def generate_transient_report(
    settings: ApiSettings,
    *,
    organization_id: str,
    city_key: str,
    report_type: str,
    output_format: str,
    visibility_tier: str = "restricted",
    title: str | None = None,
) -> dict[str, Any]:
    template = _validate_report_request(report_type, output_format)
    visibility_tier = _validate_visibility(report_type, visibility_tier)
    summary = _load_dashboard_summary(settings)
    run = _find_run(summary, city_key)
    display_name = str(run.get("display_name") or city_key.title())
    context = _build_report_context(
        settings,
        city_key,
        organization_id=organization_id,
    )
    payload, content_type, export_summary = _export_payload_for_report(report_type, output_format, context)
    report_payload = _transient_report_payload(
        organization_id=organization_id,
        city_key=city_key,
        report_type=report_type,
        output_format=output_format,
        visibility_tier=visibility_tier,
        title=title or _default_title(report_type, display_name),
        audience=str(template["audience"]),
        display_name=display_name,
    )
    return {
        "contract_version": REPORT_GENERATION_CONTRACT_VERSION,
        "persistence_mode": "transient",
        "report": report_payload,
        "artifact_access": {
            "storage_uri": None,
            "publication_tier": visibility_tier,
            "access_mode": "not_persisted",
            "download_url": "",
            "expires_in_s": None,
        },
        "artifact_summary": {
            "content_type": content_type,
            "payload_bytes": len(payload),
            **export_summary,
        },
    }


def _escape_pdf(text: str) -> str:
    return text.replace("\\", "\\\\").replace("(", "\\(").replace(")", "\\)")


# Legacy report letterhead values. The report
# serif uses the existing fallback chain (Libre Baskerville ->
# Baskerville -> Georgia/Times); Times is the closest PDF base-14 face, which
# keeps report generation dependency-free.
_PDF_INK = "0.043 0.043 0.043"
_PDF_ASPHALT = "0.063 0.075 0.090"
_PDF_NEUTRAL_500 = "0.435 0.435 0.416"
_A4_WIDTH = 595
_A4_HEIGHT = 842
_PAGE_MARGIN = 57  # ~20 mm


def _pdf_polygon(points: list[tuple[int, int]], x: float, y_top: float, scale: float) -> str:
    ops = []
    for index, (u, v) in enumerate(points):
        op = "m" if index == 0 else "l"
        ops.append(f"{x + u * scale:.2f} {y_top - v * scale:.2f} {op}")
    return "\n".join(ops) + "\nh\n"


def _pdf_brand_mark(x: float, y_top: float, size: float) -> str:
    """Render the application's Signal road-warning mark in the PDF letterhead."""
    scale = size / 64
    stream = "q\n"
    for color, points in [
        ("0.063 0.075 0.090", [(0, 0), (64, 0), (64, 64), (0, 64)]),
        ("1 0.706 0", [(32, 5), (59, 32), (32, 59), (5, 32)]),
        ("0.063 0.075 0.090", [(20, 44), (29.6, 21), (34.4, 21), (44, 44), (32, 56)]),
    ]:
        stream += f"{color} rg\n" + _pdf_polygon(points, x, y_top, scale) + "f\n"
    stream += f"1 0.706 0 RG {1.9 * scale:.2f} w\n"
    for start, end in [(52.5, 47.3), (43.2, 39.1), (35.5, 32.4), (29.1, 26.9)]:
        stream += f"{x + 32 * scale:.2f} {y_top - start * scale:.2f} m {x + 32 * scale:.2f} {y_top - end * scale:.2f} l S\n"
    # Risk-location dot from the canonical Signal icon.
    for radius, color in [(3.1, "0.063 0.075 0.090"), (1.2, "1 0.294 0.243")]:
        cx, cy, radius = x + 32 * scale, y_top - 18.4 * scale, radius * scale
        k = radius * 0.55228475
        stream += f"{color} rg {cx + radius:.2f} {cy:.2f} m\n"
        for points in [
            (cx + radius, cy + k, cx + k, cy + radius, cx, cy + radius),
            (cx - k, cy + radius, cx - radius, cy + k, cx - radius, cy),
            (cx - radius, cy - k, cx - k, cy - radius, cx, cy - radius),
            (cx + k, cy - radius, cx + radius, cy - k, cx + radius, cy),
        ]:
            stream += " ".join(f"{value:.2f}" for value in points) + " c\n"
        stream += "f\n"
    return stream + "Q\n"


def _pdf_text(x: float, y: float, font: str, size: float, colour: str, text: str) -> str:
    return f"BT {colour} rg /{font} {size} Tf {x:.2f} {y:.2f} Td ({_escape_pdf(text)}) Tj ET\n"


def _simple_pdf_bytes(title: str, lines: list[str]) -> bytes:
    """Render a corporate-letterhead PDF (Signal mark, serif titling, A4)."""
    wrapped: list[tuple[str, str]] = []  # (kind, text)
    for line in lines:
        stripped = line.strip()
        if stripped.startswith("#"):
            heading = stripped.lstrip("#").strip()
            if heading and heading.lower() != title.lower():
                wrapped.append(("gap", ""))
                for part in textwrap.wrap(heading, width=70) or [""]:
                    wrapped.append(("heading", part))
            continue
        for part in textwrap.wrap(line, width=92) or [""]:
            wrapped.append(("body", part))

    generated = datetime.now(UTC).strftime("%d %B %Y")
    line_height = 16
    first_top = _A4_HEIGHT - 200  # below letterhead + title
    cont_top = _A4_HEIGHT - 100
    bottom = _PAGE_MARGIN + 30

    pages: list[list[tuple[str, str]]] = [[]]
    cursor = first_top
    for kind, text in wrapped:
        step = line_height if kind != "gap" else 10
        if cursor - step < bottom:
            pages.append([])
            cursor = cont_top
        pages[-1].append((kind, text))
        cursor -= step

    objects: list[bytes] = []
    objects.append(b"<< /Type /Catalog /Pages 2 0 R >>")
    objects.append(b"")
    objects.append(b"<< /Type /Font /Subtype /Type1 /BaseFont /Times-Roman >>")
    objects.append(b"<< /Type /Font /Subtype /Type1 /BaseFont /Times-Bold >>")
    objects.append(b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>")

    font_resources = "<< /F1 3 0 R /F2 4 0 R /F3 5 0 R >>"
    page_refs: list[str] = []
    total = len(pages)
    for page_index, page_lines in enumerate(pages):
        page_object_id = 6 + page_index * 2
        content_object_id = page_object_id + 1
        page_refs.append(f"{page_object_id} 0 R")

        stream = ""
        rule_y = _A4_HEIGHT - 104
        if page_index == 0:
            # Letterhead: 14 mm Signal mark top-left, live-text wordmark, date top
            # right, 1 pt asphalt rule.
            mark_size = 40
            stream += _pdf_brand_mark(_PAGE_MARGIN, _A4_HEIGHT - _PAGE_MARGIN + 4, mark_size)
            stream += _pdf_text(
                _PAGE_MARGIN + mark_size + 12, _A4_HEIGHT - _PAGE_MARGIN - 12,
                "F2", 15, _PDF_INK, "xROAD",
            )
            stream += _pdf_text(
                _PAGE_MARGIN + mark_size + 12, _A4_HEIGHT - _PAGE_MARGIN - 26,
                "F3", 8, _PDF_NEUTRAL_500, "Road Accidents Risk Modeling",
            )
            stream += _pdf_text(
                _A4_WIDTH - _PAGE_MARGIN - 130, _A4_HEIGHT - _PAGE_MARGIN - 12,
                "F3", 8, _PDF_NEUTRAL_500, f"Internal - {generated}",
            )
            stream += (
                f"{_PDF_ASPHALT} RG 1 w "
                f"{_PAGE_MARGIN} {rule_y} m {_A4_WIDTH - _PAGE_MARGIN} {rule_y} l S\n"
            )
            title_y = rule_y - 40
            for part in textwrap.wrap(title, width=42)[:2]:
                stream += _pdf_text(_PAGE_MARGIN, title_y, "F2", 24, _PDF_INK, part)
                title_y -= 30
            cursor = first_top
        else:
            stream += _pdf_text(
                _PAGE_MARGIN, _A4_HEIGHT - _PAGE_MARGIN, "F2", 10, _PDF_INK, "xROAD",
            )
            stream += _pdf_text(
                _A4_WIDTH - _PAGE_MARGIN - 160, _A4_HEIGHT - _PAGE_MARGIN,
                "F3", 8, _PDF_NEUTRAL_500, title[:44],
            )
            stream += (
                f"{_PDF_ASPHALT} RG 0.5 w "
                f"{_PAGE_MARGIN} {_A4_HEIGHT - _PAGE_MARGIN - 8} m "
                f"{_A4_WIDTH - _PAGE_MARGIN} {_A4_HEIGHT - _PAGE_MARGIN - 8} l S\n"
            )
            cursor = cont_top

        for kind, text in page_lines:
            if kind == "gap":
                cursor -= 10
                continue
            if kind == "heading":
                stream += _pdf_text(_PAGE_MARGIN, cursor, "F2", 13, _PDF_INK, text)
            else:
                stream += _pdf_text(_PAGE_MARGIN, cursor, "F1", 10.5, _PDF_INK, text)
            cursor -= line_height

        footer_y = _PAGE_MARGIN - 20
        stream += (
            f"{_PDF_NEUTRAL_500} RG 0.5 w "
            f"{_PAGE_MARGIN} {footer_y + 12} m {_A4_WIDTH - _PAGE_MARGIN} {footer_y + 12} l S\n"
        )
        stream += _pdf_text(_PAGE_MARGIN, footer_y, "F3", 7.5, _PDF_NEUTRAL_500, title[:70])
        stream += _pdf_text(
            _A4_WIDTH - _PAGE_MARGIN - 60, footer_y, "F3", 7.5, _PDF_NEUTRAL_500,
            f"Page {page_index + 1} of {total}",
        )

        content_bytes = stream.encode("latin-1", errors="replace")
        objects.append(
            (
                f"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 {_A4_WIDTH} {_A4_HEIGHT}] "
                f"/Resources << /Font {font_resources} >> /Contents {content_object_id} 0 R >>"
            ).encode("latin-1")
        )
        objects.append(
            f"<< /Length {len(content_bytes)} >>\nstream\n".encode("latin-1")
            + content_bytes
            + b"\nendstream"
        )

    objects[1] = f"<< /Type /Pages /Kids [{' '.join(page_refs)}] /Count {len(page_refs)} >>".encode("latin-1")

    output = io.BytesIO()
    output.write(b"%PDF-1.4\n")
    offsets = [0]
    for index, payload in enumerate(objects, start=1):
        offsets.append(output.tell())
        output.write(f"{index} 0 obj\n".encode("latin-1"))
        output.write(payload)
        output.write(b"\nendobj\n")
    xref_offset = output.tell()
    output.write(f"xref\n0 {len(objects) + 1}\n".encode("latin-1"))
    output.write(b"0000000000 65535 f \n")
    for offset in offsets[1:]:
        output.write(f"{offset:010d} 00000 n \n".encode("latin-1"))
    output.write(
        (
            f"trailer << /Size {len(objects) + 1} /Root 1 0 R >>\n"
            f"startxref\n{xref_offset}\n%%EOF"
        ).encode("latin-1")
    )
    return output.getvalue()


def _technical_rows(brief: dict[str, Any], analyst_rows: list[dict[str, Any]]) -> list[dict[str, Any]]:
    if analyst_rows:
        return analyst_rows[:DEFAULT_SEGMENT_EXPORT_LIMIT]
    rows = []
    for item in _top_corridor_rows(brief):
        rows.append(
            {
                "segment_label": item.get("lead_segment"),
                "corridor_label": item.get("corridor_label"),
                "predicted_crash_total": item.get("predicted_crash_total"),
                "public_score_index": item.get("max_public_score_index"),
            }
        )
    return rows


def _data_quality_rows(context: dict[str, Any]) -> list[dict[str, Any]]:
    run = context["run"]
    summary = context["summary"]
    segment_scores = run.get("segment_scores") or {}
    score_tiles = run.get("score_tiles") or {}
    portfolio = context["portfolio_export"]
    signal_summary = context["executive_overview"]["reporting_context"]["signal_summary"]
    return [
        {
            "check": "published_run_count",
            "value": (summary.get("stats") or {}).get("published_run_count"),
            "status": "available" if (summary.get("stats") or {}).get("published_run_count") else "watch",
        },
        {
            "check": "segment_score_rows",
            "value": segment_scores.get("segment_score_rows"),
            "status": "available" if segment_scores.get("segment_score_rows") else "watch",
        },
        {
            "check": "public_tile_rows",
            "value": score_tiles.get("tile_rows"),
            "status": "available" if score_tiles.get("tile_rows") else "watch",
        },
        {
            "check": "managed_portfolio_rows",
            "value": portfolio.get("summary", {}).get("portfolio_count"),
            "status": portfolio.get("availability") or "unknown",
        },
        {
            "check": "signal_phase_provider_coverage",
            "value": signal_summary.get("movements_with_provider_data"),
            "status": signal_summary.get("availability") or "unknown",
        },
        {
            "check": "publication_policy",
            "value": "public outputs are restricted to aggregated score tiles and public transparency reports",
            "status": "enforced",
        },
    ]


def _public_transparency_rows(context: dict[str, Any]) -> list[dict[str, Any]]:
    executive = context["executive_overview"]
    rows = [
        {
            "section": "city_summary",
            "label": "Predicted crash burden",
            "value": executive["kpis"].get("predicted_crash_total"),
            "public_note": "Aggregated city-level model output.",
        },
        {
            "section": "city_summary",
            "label": "Very-high-risk public score tiles",
            "value": executive["kpis"].get("very_high_risk_tiles"),
            "public_note": "Aggregated public tile count.",
        },
        {
            "section": "city_summary",
            "label": "Completed methodology context",
            "value": executive["kpis"].get("model_confidence"),
            "public_note": "Confidence label derived from transferability checks.",
        },
    ]
    for row in executive["priority_areas"]["top_corridors"]:
        rows.append(
            {
                "section": "priority_corridor",
                "label": row.get("corridor_label"),
                "value": row.get("max_public_score_index"),
                "public_note": "Public score index only; no restricted crash record payloads.",
            }
        )
    return rows


def _build_report_context(
    settings: ApiSettings,
    city_key: str,
    *,
    session: Any | None = None,
    organization_id: str | None = None,
) -> dict[str, Any]:
    summary = _load_dashboard_summary(settings)
    run = _find_run(summary, city_key)
    brief = _load_intervention_brief(settings, run)
    analyst_csv_rows = _load_csv_rows(_segment_csv_path(settings, run), limit=DEFAULT_SEGMENT_EXPORT_LIMIT)
    segment_geojson = _load_geojson(_segment_geojson_path(settings, run))
    tile_geojson = _load_geojson(_tile_geojson_path(settings, run))
    portfolio_export = (
        intervention_portfolio.build_portfolio_export(
            session,
            organization_id=organization_id,
            city_key=city_key,
        )
        if session is not None and organization_id is not None
        else _empty_portfolio_export(organization_id, city_key)
    )
    return {
        "summary": summary,
        "run": run,
        "brief": brief,
        "executive_overview": build_executive_overview(settings, city_key=city_key),
        "analyst_rows": analyst_csv_rows,
        "segment_geojson": segment_geojson,
        "tile_geojson": tile_geojson,
        "portfolio_export": portfolio_export,
    }


def _render_markdown(report_type: str, context: dict[str, Any]) -> str:
    if report_type == "intervention_portfolio":
        portfolio = context["portfolio_export"]
        lines = [
            "# Intervention Portfolio",
            "",
            "## Summary",
            f"- Portfolio rows: `{portfolio['summary']['portfolio_count']}`",
            f"- Expected benefit: `{portfolio['summary']['expected_benefit']}`",
            "",
            "## Portfolio Rows",
        ]
        for row in portfolio["rows"]:
            lines.append(
                f"- P{row.get('priority_rank')}: {row.get('target')} · {row.get('intervention_type')} · "
                f"{row.get('status')} · benefit `{row.get('expected_benefit')}`"
            )
        if not portfolio["rows"]:
            lines.append("- No managed portfolio rows available.")
        return "\n".join(lines) + "\n"

    if report_type == "data_quality_governance":
        city = context["executive_overview"]["city"]["display_name"]
        lines = [
            f"# {city} Data Quality & Governance Report",
            "",
            "## Production Readiness Checks",
        ]
        for row in _data_quality_rows(context):
            lines.append(f"- {row['check']}: `{row['value']}` ({row['status']})")
        lines.extend(
            [
                "",
                "## Governance",
                "- Public outputs are limited to aggregated score tiles, public corridors, and approved transparency reports.",
                "- Restricted technical, governance, and portfolio reports require authenticated organization access.",
            ]
        )
        return "\n".join(lines) + "\n"

    if report_type == "public_transparency":
        executive = context["executive_overview"]
        city = executive["city"]["display_name"]
        lines = [
            f"# {city} Public Transparency Report",
            "",
            "## Public Safety Summary",
            f"- Predicted crash burden: `{executive['kpis']['predicted_crash_total']}`",
            f"- Very-high-risk public score tiles: `{executive['kpis']['very_high_risk_tiles']}`",
            f"- Model confidence: `{executive['kpis']['model_confidence']}`",
            "",
            "## Public Priority Corridors",
        ]
        for row in executive["priority_areas"]["top_corridors"]:
            lines.append(
                f"- {row.get('corridor_label')}: public score `{row.get('max_public_score_index')}`"
            )
        lines.extend(
            [
                "",
                "## Data Protection",
                "- This public report excludes raw crash records, restricted analyst exports, and private source payloads.",
            ]
        )
        return "\n".join(lines) + "\n"

    executive = context["executive_overview"]
    city = executive["city"]["display_name"]
    kpis = executive["kpis"]
    budget = executive["budget_impact"]
    lines = [
        f"# {city} {REPORT_TEMPLATES[report_type]['label']}",
        "",
        "## KPI",
        f"- Predicted crash burden: `{kpis['predicted_crash_total']}`",
        f"- Observed crash total: `{kpis['observed_crash_total']}`",
        f"- Very-high-risk tiles: `{kpis['very_high_risk_tiles']}`",
        f"- High-risk corridors: `{kpis['high_risk_corridor_count']}`",
        f"- Model confidence: `{kpis['model_confidence']}`",
        "",
        "## Recommended portfolio",
        f"- Title: {budget['recommended_portfolio_title'] or 'n/a'}",
        f"- Expected crash reduction: `{budget['expected_crash_reduction']}`",
        f"- Expected relative reduction: `{budget['expected_relative_reduction']}`",
        "",
        "## Priority areas",
    ]
    for row in executive["priority_areas"]["top_corridors"]:
        lines.append(
            f"- {row.get('corridor_label')}: predicted crashes `{row.get('predicted_crash_total')}`, "
            f"score `{row.get('max_public_score_index')}`"
        )
    if report_type == "technical_hotspot":
        lines.extend(
            [
                "",
                "## Recommended actions",
            ]
        )
        for row in executive["priority_areas"]["top_actions"]:
            lines.append(f"- P{row.get('priority')}: {row.get('title')}")
    return "\n".join(lines) + "\n"


def _render_csv(rows: list[dict[str, Any]]) -> bytes:
    output = io.StringIO()
    fieldnames = list(rows[0].keys()) if rows else ["note"]
    writer = csv.DictWriter(output, fieldnames=fieldnames)
    writer.writeheader()
    if rows:
        writer.writerows(rows)
    else:
        writer.writerow({"note": "No rows available"})
    return output.getvalue().encode("utf-8")


def _render_xlsx(report_type: str, context: dict[str, Any]) -> bytes:
    from openpyxl import Workbook

    if report_type == "data_quality_governance":
        workbook = Workbook()
        quality_sheet = workbook.active
        quality_sheet.title = "quality"
        quality_sheet.append(["check", "value", "status"])
        for row in _data_quality_rows(context):
            quality_sheet.append([row.get("check"), row.get("value"), row.get("status")])
        governance_sheet = workbook.create_sheet("governance")
        governance_sheet.append(["policy", "status"])
        governance_sheet.append(["public aggregation only", "enforced"])
        governance_sheet.append(["restricted downloads require role policy", "enforced"])
        buffer = io.BytesIO()
        workbook.save(buffer)
        return buffer.getvalue()

    if report_type == "public_transparency":
        workbook = Workbook()
        overview_sheet = workbook.active
        overview_sheet.title = "public_summary"
        rows = _public_transparency_rows(context)
        overview_sheet.append(["section", "label", "value", "public_note"])
        for row in rows:
            overview_sheet.append([row.get("section"), row.get("label"), row.get("value"), row.get("public_note")])
        buffer = io.BytesIO()
        workbook.save(buffer)
        return buffer.getvalue()

    if report_type == "intervention_portfolio":
        workbook = Workbook()
        portfolio_sheet = workbook.active
        portfolio_sheet.title = "portfolio"
        rows = context["portfolio_export"]["rows"]
        fieldnames = list(rows[0].keys()) if rows else ["note"]
        portfolio_sheet.append(fieldnames)
        if rows:
            for row in rows:
                portfolio_sheet.append([row.get(name) for name in fieldnames])
        else:
            portfolio_sheet.append(["No managed portfolio rows available"])
        buffer = io.BytesIO()
        workbook.save(buffer)
        return buffer.getvalue()

    executive = context["executive_overview"]
    workbook = Workbook()
    overview_sheet = workbook.active
    overview_sheet.title = "overview"
    overview_sheet.append(["KPI", "Value"])
    for key, value in executive["kpis"].items():
        overview_sheet.append([key, value])

    corridor_sheet = workbook.create_sheet("priority_areas")
    corridor_sheet.append(["corridor_label", "predicted_crash_total", "max_public_score_index", "segment_count", "lead_segment"])
    for row in executive["priority_areas"]["top_corridors"]:
        corridor_sheet.append(
            [
                row.get("corridor_label"),
                row.get("predicted_crash_total"),
                row.get("max_public_score_index"),
                row.get("segment_count"),
                row.get("lead_segment"),
            ]
        )

    if report_type == "technical_hotspot":
        technical_sheet = workbook.create_sheet("technical_rows")
        rows = _technical_rows(context["brief"], context["analyst_rows"])
        fieldnames = list(rows[0].keys()) if rows else ["note"]
        technical_sheet.append(fieldnames)
        for row in rows:
            technical_sheet.append([row.get(name) for name in fieldnames])

    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def _export_payload_for_report(report_type: str, output_format: str, context: dict[str, Any]) -> tuple[bytes, str, dict[str, Any]]:
    executive = context["executive_overview"]
    if output_format == "md":
        markdown = _render_markdown(report_type, context)
        return markdown.encode("utf-8"), "text/markdown", {"line_count": len(markdown.splitlines())}
    if output_format == "pdf":
        markdown_lines = _render_markdown(report_type, context).splitlines()
        pdf_bytes = _simple_pdf_bytes(executive["city"]["display_name"], markdown_lines)
        return pdf_bytes, "application/pdf", {"page_hint": max(1, len(markdown_lines) // 42 + 1)}
    if output_format == "csv":
        if report_type == "intervention_portfolio":
            rows = context["portfolio_export"]["rows"]
        elif report_type == "data_quality_governance":
            rows = _data_quality_rows(context)
        elif report_type == "public_transparency":
            rows = _public_transparency_rows(context)
        else:
            rows = _technical_rows(context["brief"], context["analyst_rows"])
        return _render_csv(rows), "text/csv", {"row_count": len(rows)}
    if output_format == "geojson":
        geojson = context["segment_geojson"] if report_type == "technical_hotspot" else context["tile_geojson"]
        return (
            json.dumps(geojson, indent=2, sort_keys=True).encode("utf-8"),
            "application/geo+json",
            {"feature_count": len(geojson.get("features", []))},
        )
    if output_format == "xlsx":
        xlsx_bytes = _render_xlsx(report_type, context)
        if report_type == "intervention_portfolio":
            return (
                xlsx_bytes,
                "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                {"sheet_count": 1, "row_count": len(context["portfolio_export"]["rows"])},
            )
        if report_type == "data_quality_governance":
            return (
                xlsx_bytes,
                "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                {"sheet_count": 2, "row_count": len(_data_quality_rows(context))},
            )
        if report_type == "public_transparency":
            return (
                xlsx_bytes,
                "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                {"sheet_count": 1, "row_count": len(_public_transparency_rows(context))},
            )
        return (
            xlsx_bytes,
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            {"sheet_count": 3 if report_type == "technical_hotspot" else 2},
        )
    raise ReportingHubError(f"Unsupported output_format: {output_format}")


def execute_report_generation(session: Any, settings: ApiSettings, report: Any) -> dict[str, Any]:
    report_type = str(getattr(report, "report_type", "") or "")
    output_format = str(getattr(report, "output_format", "") or "")
    parameters = dict(getattr(report, "parameters_json", {}) or {})
    city_key = str(parameters.get("city_key") or "").strip()
    if not city_key:
        raise ReportingHubError("Report is missing city_key in parameters_json.")

    _validate_report_request(report_type, output_format)
    report.status = "running"
    if hasattr(session, "flush"):
        session.flush()

    context = _build_report_context(
        settings,
        city_key,
        session=session,
        organization_id=str(getattr(report, "organization_id")),
    )
    payload, content_type, export_summary = _export_payload_for_report(report_type, output_format, context)
    file_stem = city_key.replace("/", "_")
    filename = f"{file_stem}_{report_type}.{output_format}"

    persisted = artifact_storage.persist_generated_artifact(
        session,
        organization_id=str(getattr(report, "organization_id")),
        municipality_id=None,
        artifact_kind=f"{report_type}_report_output",
        artifact_name=str(getattr(report, "title", report_type)),
        filename=filename,
        payload=payload,
        report=report,
        publication_tier=str(getattr(report, "visibility_tier", "restricted")),
        content_type=content_type,
        linked_objects={"city_key": city_key, "report_type": report_type, "output_format": output_format},
        lineage={
            "reporting_contract_version": REPORT_GENERATION_CONTRACT_VERSION,
            "city_key": city_key,
            "report_type": report_type,
            "output_format": output_format,
        },
    )
    summary = dict(getattr(report, "summary_json", {}) or {})
    summary.update(
        {
            "export_summary": export_summary,
            "review_state": "pending_review",
            "generated_contract_version": REPORT_GENERATION_CONTRACT_VERSION,
        }
    )
    report.summary_json = summary
    if hasattr(session, "flush"):
        session.flush()
    repositories.record_audit_log(
        session,
        organization_id=str(getattr(report, "organization_id")),
        action="report.generated",
        target_type="report",
        target_id=str(getattr(report, "id", "")),
        status="generated",
        metadata={
            "city_key": city_key,
            "report_type": report_type,
            "output_format": output_format,
            **export_summary,
        },
    )
    return {
        "contract_version": REPORT_GENERATION_CONTRACT_VERSION,
        "report": _serialize_report(report),
        "artifact_access": persisted["access"].as_dict(),
    }


def _has_role(role_names: set[str], required_role: str) -> bool:
    if not required_role:
        return True
    return required_role in role_names


def _required_download_role(report: Any) -> str:
    visibility_tier = str(getattr(report, "visibility_tier", "restricted") or "restricted")
    report_type = str(getattr(report, "report_type", "") or "")
    if visibility_tier == "analyst_only":
        return "analyst"
    template_role = required_role_for_report(report_type)
    if visibility_tier in {"restricted", "internal", "partner_restricted"} and template_role == "municipality_viewer":
        return "municipality_viewer"
    return template_role


def _assert_download_allowed(
    report: Any,
    *,
    role_names: set[str] | None = None,
    public_request: bool = False,
) -> None:
    status = str(getattr(report, "status", "") or "")
    visibility_tier = str(getattr(report, "visibility_tier", "restricted") or "restricted")
    if public_request:
        if status != "published" or visibility_tier != "public":
            raise ReportingHubAuthorizationError("Public report download requires published public visibility.")
        return
    if status not in REPORT_DOWNLOAD_READY_STATUSES:
        raise ReportingHubError("Report download is not ready yet.")
    if role_names is None:
        return
    required_role = _required_download_role(report)
    if not _has_role(role_names, required_role):
        raise ReportingHubAuthorizationError(f"Report download requires role: {required_role}.")


def _lifecycle_summary(action: str, actor_subject: str | None, note: str | None = None) -> dict[str, Any]:
    return {
        "review_state": {
            "review": "reviewed",
            "publish": "published",
            "unpublish": "reviewed",
            "archive": "archived",
            "reopen": "pending_review",
            "fail": "failed",
        }.get(action, action),
        "last_lifecycle_action": action,
        "last_lifecycle_actor": actor_subject,
        "last_lifecycle_note": note,
        "last_lifecycle_at": datetime.now(UTC).isoformat(),
    }


def transition_report_lifecycle(
    session: Any,
    *,
    organization_id: str,
    report_id: str,
    action: str,
    actor_subject: str | None = None,
    visibility_tier: str | None = None,
    note: str | None = None,
) -> dict[str, Any]:
    report = repositories.get_report(session, report_id)
    if report is None or str(getattr(report, "organization_id", "")) != organization_id:
        raise ReportingHubError("Report not found.")

    action = (action or "").strip().lower()
    current_status = str(getattr(report, "status", "") or "")
    if action == "review":
        if current_status != "generated":
            raise ReportingHubError("Only generated reports can be marked reviewed.")
        repositories.update_report_lifecycle(
            session,
            report,
            status="reviewed",
            summary=_lifecycle_summary(action, actor_subject, note),
        )
    elif action == "publish":
        if current_status not in {"reviewed", "published"}:
            raise ReportingHubError("Only reviewed reports can be published.")
        if not getattr(report, "file_storage_uri", None):
            raise ReportingHubError("Report must have a generated artifact before publishing.")
        tier = _validate_visibility(str(getattr(report, "report_type", "")), visibility_tier or "public")
        repositories.update_report_lifecycle(
            session,
            report,
            status="published",
            visibility_tier=tier,
            published_at=datetime.now(UTC),
            summary=_lifecycle_summary(action, actor_subject, note),
        )
    elif action == "unpublish":
        if current_status != "published":
            raise ReportingHubError("Only published reports can be unpublished.")
        repositories.update_report_lifecycle(
            session,
            report,
            status="reviewed",
            visibility_tier=visibility_tier or "restricted",
            clear_published_at=True,
            summary=_lifecycle_summary(action, actor_subject, note),
        )
    elif action == "archive":
        if current_status in {"draft", "queued", "running"}:
            raise ReportingHubError("Only completed report records can be archived.")
        repositories.update_report_lifecycle(
            session,
            report,
            status="archived",
            summary=_lifecycle_summary(action, actor_subject, note),
        )
    elif action == "reopen":
        if current_status not in {"archived", "failed"}:
            raise ReportingHubError("Only archived or failed reports can be reopened.")
        repositories.update_report_lifecycle(
            session,
            report,
            status="generated" if getattr(report, "file_storage_uri", None) else "draft",
            clear_published_at=True,
            summary=_lifecycle_summary(action, actor_subject, note),
        )
    elif action == "fail":
        repositories.update_report_lifecycle(
            session,
            report,
            status="failed",
            summary=_lifecycle_summary(action, actor_subject, note),
        )
    else:
        raise ReportingHubError("Unsupported report lifecycle action.")

    repositories.record_audit_log(
        session,
        organization_id=organization_id,
        action=f"report.{action}",
        target_type="report",
        target_id=str(getattr(report, "id", "")),
        status=str(getattr(report, "status", "")),
        metadata={
            "report_type": getattr(report, "report_type", None),
            "visibility_tier": getattr(report, "visibility_tier", None),
            "note": note,
            "actor_subject": actor_subject,
        },
    )
    return {
        "contract_version": REPORT_LIFECYCLE_CONTRACT_VERSION,
        "report": _serialize_report(report),
    }


def get_report_detail(
    session: Any,
    *,
    organization_id: str,
    report_id: str,
    role: str = REPORT_DEFAULT_AUTHENTICATED_ROLE,
) -> dict[str, Any]:
    """Return one report's matrix-projected detail block.

    `role` controls the matrix projection just like `list_report_history`;
    defaults to `"analyst"`. The response includes the `redaction_policy`
    audit block.
    """
    report = repositories.get_report(session, report_id)
    if report is None or str(getattr(report, "organization_id", "")) != organization_id:
        raise ReportingHubError("Report not found.")
    return {
        "contract_version": REPORT_DETAIL_CONTRACT_VERSION,
        "report": _serialize_report_for_role(report, role),
        "redaction_policy": _report_redaction_policy_meta(role),
    }


def get_report_download(
    session: Any,
    *,
    organization_id: str,
    report_id: str,
    role_names: set[str] | None = None,
) -> dict[str, Any]:
    report = repositories.get_report(session, report_id)
    if report is None or str(getattr(report, "organization_id", "")) != organization_id:
        raise ReportingHubError("Report not found.")
    if not getattr(report, "file_storage_uri", None):
        raise ReportingHubError("Report download is not ready yet.")
    _assert_download_allowed(report, role_names=role_names)
    access = resolve_artifact_download(
        str(report.file_storage_uri),
        publication_tier=str(getattr(report, "visibility_tier", "restricted")),
    )
    return {
        "contract_version": REPORT_DETAIL_CONTRACT_VERSION,
        "report": _serialize_report(report),
        "access": access.as_dict(),
    }


def list_public_reports(
    session: Any,
    *,
    city_key: str,
    limit: int = DEFAULT_HISTORY_LIMIT,
) -> dict[str, Any]:
    """List publicly-published reports for a city, matrix-projected for `"public"`.

    The matrix entries for `report` with `"public"` in their
    `role_visibility` are the only fields surfaced. Notably,
    `organization_id`, `audience`, `summary`, `review_state`,
    `lifecycle_actions`, `created_at`, and `file_storage_uri` are
    default-denied for this surface. The response includes a
    `redaction_policy` audit block.
    """
    try:
        reports = repositories.list_public_reports_for_city(session, city_key, limit=limit)
        availability = "available"
        unavailable_reason = None
    except (SQLAlchemyError, AttributeError):
        reports = []
        availability = "degraded"
        unavailable_reason = "database_unavailable"
    return {
        "contract_version": PUBLIC_REPORTS_CONTRACT_VERSION,
        "city_key": city_key,
        "availability": availability,
        "unavailable_reason": unavailable_reason,
        "rows": [_public_report_row(report) for report in reports],
        "redaction_policy": _report_redaction_policy_meta(REPORT_PUBLIC_ROLE),
    }


def get_public_report_download(session: Any, *, city_key: str, report_id: str) -> dict[str, Any]:
    report = repositories.get_report(session, report_id)
    if report is None:
        raise ReportingHubError("Report not found.")
    parameters = dict(getattr(report, "parameters_json", {}) or {})
    if str(parameters.get("city_key") or "") != city_key:
        raise ReportingHubError("Report not found.")
    if not getattr(report, "file_storage_uri", None):
        raise ReportingHubError("Report download is not ready yet.")
    _assert_download_allowed(report, public_request=True)
    access = resolve_artifact_download(
        str(report.file_storage_uri),
        publication_tier=str(getattr(report, "visibility_tier", "public")),
    )
    return {
        "contract_version": REPORT_DETAIL_CONTRACT_VERSION,
        "report": _public_report_row(report),
        "access": access.as_dict(),
    }


__all__ = [
    "DEFAULT_HISTORY_LIMIT",
    "EXECUTIVE_OVERVIEW_CONTRACT_VERSION",
    "REPORT_DETAIL_CONTRACT_VERSION",
    "REPORT_GENERATION_CONTRACT_VERSION",
    "REPORT_GENERATION_TASK_NAME",
    "REPORT_HISTORY_CONTRACT_VERSION",
    "REPORT_LIFECYCLE_CONTRACT_VERSION",
    "PUBLIC_REPORTS_CONTRACT_VERSION",
    "REPORT_TEMPLATES",
    "REPORTING_OVERVIEW_CONTRACT_VERSION",
    "ReportingHubAuthorizationError",
    "ReportingHubError",
    "build_executive_overview",
    "build_reporting_overview",
    "create_report_request",
    "execute_report_generation",
    "generate_transient_report",
    "get_public_report_download",
    "get_report_detail",
    "get_report_download",
    "list_public_reports",
    "list_report_history",
    "required_role_for_report",
    "transition_report_lifecycle",
]
